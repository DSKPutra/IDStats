"""Manajemen data: baca CSV/XLSX, pembersihan (missing value, outlier), dan transformasi.

Sengaja tanpa pandas agar bundle fungsi serverless tetap kecil (pandas ±40 MB): kolom numerik disimpan
sebagai array NumPy (NaN = kosong), kolom teks sebagai list.
"""

from __future__ import annotations

import csv
import io
from typing import Any

import numpy as np

from app.core.errors import SolverError
from app.schemas.common import NamedTable, SolverResponse, Step, Table
from app.solvers._base import fmt, summary_items

MAX_ROWS = 10_000
MAX_BYTES = 5 * 1024 * 1024


class Frame:
    """Tabel data minimal: urutan kolom + isi per kolom."""

    def __init__(self, columns: list[str], cols: dict[str, Any]):
        self.columns = columns
        self.cols = cols

    def __len__(self) -> int:
        return len(next(iter(self.cols.values()))) if self.cols else 0

    def is_numeric(self, c: str) -> bool:
        return isinstance(self.cols[c], np.ndarray)

    def numeric_columns(self) -> list[str]:
        return [c for c in self.columns if self.is_numeric(c)]

    def missing(self, c: str) -> np.ndarray:
        v = self.cols[c]
        return np.isnan(v) if isinstance(v, np.ndarray) else np.array([x is None for x in v], dtype=bool)

    def take(self, mask: np.ndarray) -> Frame:
        return Frame(
            list(self.columns),
            {
                c: (v[mask] if isinstance(v, np.ndarray) else [x for x, keep in zip(v, mask, strict=True) if keep])
                for c, v in self.cols.items()
            },
        )

    def rows(self) -> list[list[Any]]:
        out = []
        for i in range(len(self)):
            row = []
            for c in self.columns:
                v = self.cols[c][i]
                if isinstance(v, np.bool_ | bool):
                    row.append(bool(v))
                elif isinstance(v, float | np.floating):
                    row.append(
                        None if np.isnan(v) else (int(v) if float(v).is_integer() and abs(v) < 1e15 else float(v))
                    )
                else:
                    row.append(v)
            out.append(row)
        return out


def _to_number(v: Any, comma_decimal: bool = False) -> float | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, int | float):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    if comma_decimal:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan  # penanda gagal


def _build(columns: list[str], raw_rows: list[list[Any]], comma_decimal: bool = False) -> Frame:
    cols: dict[str, Any] = {}
    for j, c in enumerate(columns):
        values = [r[j] if j < len(r) else None for r in raw_rows]
        values = [None if (isinstance(v, str) and v.strip() == "") else v for v in values]
        nums = [_to_number(v, comma_decimal) for v in values]
        filled = [n for n, v in zip(nums, values, strict=True) if v is not None]
        if filled and all(n is not None and not (isinstance(n, float) and np.isnan(n)) for n in filled):
            cols[c] = np.array([np.nan if n is None else n for n in nums], dtype=float)
        elif all(isinstance(v, bool) or v is None for v in values) and any(v is not None for v in values):
            cols[c] = values
        else:
            cols[c] = [None if v is None else (v if isinstance(v, bool) else str(v).strip()) for v in values]
    return Frame(list(columns), cols)


def _frame(columns: list[str], rows: list[list[Any]]) -> Frame:
    if not columns:
        raise SolverError("Dataset belum memiliki kolom.")
    if len(set(columns)) != len(columns):
        raise SolverError("Nama kolom harus unik.")
    if any(len(r) != len(columns) for r in rows):
        raise SolverError("Setiap baris harus memiliki jumlah sel yang sama dengan jumlah kolom.")
    return _build(columns, rows)


def _profile(df: Frame) -> NamedTable:
    rows = []
    for c in df.columns:
        miss = df.missing(c)
        numeric = df.is_numeric(c)
        rows.append(
            [
                c,
                "Numerik" if numeric else "Teks/kategori",
                int((~miss).sum()),
                int(miss.sum()),
                float(np.nanmean(df.cols[c])) if numeric and (~miss).any() else None,
            ]
        )
    return NamedTable(title="Profil kolom", columns=["Kolom", "Tipe", "Terisi", "Kosong", "Rata-rata"], rows=rows)


def _dataset_response(df: Frame, steps: list[Step], warnings: list[str], extra: dict | None = None) -> SolverResponse:
    return SolverResponse(
        result={"columns": list(df.columns), "rows": df.rows(), "n_rows": len(df), **(extra or {})},
        steps=steps,
        tables=[_profile(df)],
        summary=summary_items(
            [("Jumlah baris", len(df)), ("Jumlah kolom", len(df.columns)), ("Kolom numerik", len(df.numeric_columns()))]
        ),
        warnings=warnings,
    )


def _read_csv(text: str) -> tuple[list[str], list[list[Any]]]:
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    data = [r for r in reader if any(cell.strip() for cell in r)]
    if not data:
        raise SolverError("File tidak berisi data.")
    header, body = data[0], data[1:]
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    header = header + [""] * (width - len(header))
    return header, [r + [""] * (width - len(r)) for r in body]


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    data = [
        list(r) for r in ws.iter_rows(values_only=True) if r and any(v is not None and str(v).strip() != "" for v in r)
    ]
    wb.close()
    if not data:
        raise SolverError("Lembar kerja pertama kosong.")
    header = ["" if v is None else str(v) for v in data[0]]
    return header, data[1:]


def read_file(content: bytes, filename: str) -> SolverResponse:
    if len(content) > MAX_BYTES:
        raise SolverError("Ukuran file maksimal 5 MB.")
    if not content:
        raise SolverError("File kosong.")
    name = filename.lower()
    try:
        if name.endswith(".xlsx"):
            header, body = _read_xlsx(content)
        elif name.endswith(".xls"):
            raise SolverError("Format .xls lama tidak didukung; simpan ulang sebagai .xlsx atau CSV.")
        elif name.endswith((".csv", ".txt", ".tsv")):
            header, body = _read_csv(content.decode("utf-8-sig", errors="replace"))
        else:
            raise SolverError("Format file harus CSV atau XLSX.")
    except SolverError:
        raise
    except Exception as exc:  # noqa: BLE001 - pesan ramah untuk file rusak/aneh
        raise SolverError(f"File tidak dapat dibaca: {exc}") from exc
    warnings = []
    if len(body) > MAX_ROWS:
        warnings.append(f"File berisi {len(body)} baris; hanya {MAX_ROWS} baris pertama yang dimuat.")
        body = body[:MAX_ROWS]
    columns = []
    for i, c in enumerate(header):
        name_c = str(c).strip() or f"Kolom{i + 1}"
        while name_c in columns:
            name_c += "_"
        columns.append(name_c)
    df = _build(columns, body)
    # Excel/CSV lokal Indonesia sering memakai koma desimal ("72,5").
    retry = _build(columns, body, comma_decimal=True)
    for c in columns:
        if not df.is_numeric(c) and retry.is_numeric(c):
            df.cols[c] = retry.cols[c]
    steps = [
        Step(
            title="Baca file",
            explanation=f"File “{filename}” dibaca: {len(df)} baris × {len(columns)} kolom. Pemisah kolom dideteksi otomatis.",
        )
    ]
    return _dataset_response(df, steps, warnings)


def _quantile(x: np.ndarray, q: float) -> float:
    return float(np.nanpercentile(x, q * 100))


def _mode(x: np.ndarray) -> float:
    vals, counts = np.unique(x[~np.isnan(x)], return_counts=True)
    return float(vals[np.argmax(counts)])


def clean(
    columns: list[str],
    rows: list[list[Any]],
    missing: str = "none",
    outlier: str = "none",
    outlier_action: str = "flag",
    iqr_k: float = 1.5,
    z_threshold: float = 3.0,
    target_columns: list[str] | None = None,
) -> SolverResponse:
    df = _frame(columns, rows)
    numeric = df.numeric_columns()
    targets = [c for c in (target_columns or numeric) if c in df.columns]
    non_numeric = [c for c in targets if c not in numeric]
    if non_numeric:
        raise SolverError(f"Kolom berikut bukan numerik: {', '.join(non_numeric)}.")
    steps: list[Step] = []
    warnings: list[str] = []
    report: list[list[Any]] = []

    if outlier != "none":
        if outlier not in ("iqr", "zscore"):
            raise SolverError("Metode outlier harus 'iqr' atau 'zscore'.")
        if outlier_action not in ("flag", "remove", "winsorize"):
            raise SolverError("Tindakan outlier harus 'flag', 'remove', atau 'winsorize'.")
        mask_any = np.zeros(len(df), dtype=bool)
        rule = ""
        for c in targets:
            col = df.cols[c]
            if outlier == "iqr":
                q1, q3 = _quantile(col, 0.25), _quantile(col, 0.75)
                lo, hi = q1 - iqr_k * (q3 - q1), q3 + iqr_k * (q3 - q1)
                rule = rf"[Q_1 - {fmt(iqr_k)}\,IQR,\ Q_3 + {fmt(iqr_k)}\,IQR]"
            else:
                mu, sd = float(np.nanmean(col)), float(np.nanstd(col, ddof=1))
                lo, hi = mu - z_threshold * sd, mu + z_threshold * sd
                rule = rf"[\bar{{x}} - {fmt(z_threshold)}s,\ \bar{{x}} + {fmt(z_threshold)}s]"
            with np.errstate(invalid="ignore"):
                mask = (col < lo) | (col > hi)
            report.append([c, float(lo), float(hi), int(mask.sum())])
            if outlier_action == "winsorize":
                df.cols[c] = np.where(np.isnan(col), np.nan, np.clip(col, lo, hi))
            mask_any |= mask
        steps.append(
            Step(
                title="Deteksi outlier",
                explanation="Data di luar batas dianggap outlier. Batas dihitung dari nilai yang terisi, sebelum imputasi.",
                latex=rf"\text{{Batas wajar: }} {rule}",
                table=Table(columns=["Kolom", "Batas bawah", "Batas atas", "Jumlah outlier"], rows=report),
            )
        )
        n_out = int(mask_any.sum())
        if outlier_action == "remove":
            df = df.take(~mask_any)
            steps.append(Step(title="Hapus outlier", explanation=f"{n_out} baris yang mengandung outlier dihapus."))
        elif outlier_action == "winsorize":
            steps.append(Step(title="Winsorisasi", explanation="Nilai outlier dipotong ke batas bawah/atas terdekat."))
        else:
            df.columns.append("outlier")
            df.cols["outlier"] = [bool(x) for x in mask_any]
            steps.append(Step(title="Tandai outlier", explanation=f"Kolom baru “outlier” = True pada {n_out} baris."))
    # Imputasi setelah penanganan outlier agar nilai pengisi tidak tertarik oleh outlier.
    before = len(df)
    if missing == "drop":
        keep = np.ones(len(df), dtype=bool)
        for c in targets:
            keep &= ~df.missing(c)
        df = df.take(keep)
        steps.append(
            Step(
                title="Hapus baris dengan nilai kosong",
                explanation=f"{before - len(df)} baris dihapus (listwise deletion).",
            )
        )
    elif missing in ("mean", "median", "mode"):
        fills = []
        for c in targets:
            col = df.cols[c]
            n_missing = int(np.isnan(col).sum())
            if not n_missing:
                continue
            if np.isnan(col).all():
                raise SolverError(f"Kolom {c} seluruhnya kosong sehingga tidak dapat diimputasi.")
            value = {"mean": float(np.nanmean(col)), "median": float(np.nanmedian(col)), "mode": _mode(col)}[missing]
            df.cols[c] = np.where(np.isnan(col), value, col)
            fills.append([c, n_missing, value])
        label = {"mean": "rata-rata", "median": "median", "mode": "modus"}[missing]
        steps.append(
            Step(
                title=f"Imputasi nilai kosong dengan {label}",
                explanation="Tidak ada nilai kosong pada kolom terpilih."
                if not fills
                else f"Nilai kosong diganti {label} kolomnya.",
                table=Table(columns=["Kolom", "Jumlah diisi", "Nilai pengisi"], rows=fills) if fills else None,
            )
        )
    elif missing != "none":
        raise SolverError("Metode missing value tidak dikenal.")

    if not steps:
        warnings.append("Tidak ada operasi pembersihan yang dipilih.")
    return _dataset_response(df, steps, warnings, {"outliers": report})


def transform(columns: list[str], rows: list[list[Any]], targets: list[str], method: str) -> SolverResponse:
    df = _frame(columns, rows)
    if not targets:
        raise SolverError("Pilih minimal satu kolom untuk ditransformasi.")
    numeric = df.numeric_columns()
    formulas = {
        "zscore": (r"z = \frac{x - \bar{x}}{s}", "z"),
        "minmax": (r"x' = \frac{x - x_{min}}{x_{max} - x_{min}}", "minmax"),
        "log": (r"x' = \ln x", "ln"),
        "log10": (r"x' = \log_{10} x", "log10"),
        "sqrt": (r"x' = \sqrt{x}", "sqrt"),
    }
    if method not in formulas:
        raise SolverError("Metode transformasi tidak dikenal.")
    formula, suffix = formulas[method]
    info = []
    for c in targets:
        if c not in numeric:
            raise SolverError(f"Kolom {c} bukan numerik.")
        x = df.cols[c]
        with np.errstate(invalid="ignore", divide="ignore"):
            if method == "zscore":
                sd = float(np.nanstd(x, ddof=1))
                if not sd:
                    raise SolverError(f"Kolom {c} konstan (s = 0).")
                new = (x - np.nanmean(x)) / sd
                info.append([c, f"x̄ = {fmt(float(np.nanmean(x)))}, s = {fmt(sd)}"])
            elif method == "minmax":
                lo, hi = float(np.nanmin(x)), float(np.nanmax(x))
                if hi == lo:
                    raise SolverError(f"Kolom {c} konstan (max = min).")
                new = (x - lo) / (hi - lo)
                info.append([c, f"min = {fmt(lo)}, max = {fmt(hi)}"])
            elif method in ("log", "log10"):
                if np.any(x[~np.isnan(x)] <= 0):
                    raise SolverError(f"Kolom {c} memuat nilai ≤ 0; logaritma hanya untuk nilai positif.")
                new = np.log(x) if method == "log" else np.log10(x)
                info.append([c, "semua nilai positif"])
            else:
                if np.any(x[~np.isnan(x)] < 0):
                    raise SolverError(f"Kolom {c} memuat nilai negatif; akar kuadrat hanya untuk nilai ≥ 0.")
                new = np.sqrt(x)
                info.append([c, "semua nilai ≥ 0"])
        name = f"{c}_{suffix}"
        df.columns.append(name)
        df.cols[name] = new
    steps = [
        Step(
            title="Terapkan transformasi",
            explanation="Kolom hasil ditambahkan di sebelah kanan dengan akhiran _" + suffix + ".",
            latex=formula,
            table=Table(columns=["Kolom", "Parameter"], rows=info),
        )
    ]
    return _dataset_response(df, steps, [])
